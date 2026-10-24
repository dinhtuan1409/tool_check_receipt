"""Xuất file theo mẫu nhập phần mềm hóa đơn (V7): cấu trúc khối hóa đơn, dòng giảm giá, sheet Kiểm tra."""
from __future__ import annotations

import io

import openpyxl
import pandas as pd
import pytest

import invoice_engine as eng
from test_engine_edge_cases import (COLD_BASIC, HOT_BASIC, catalog_from, load_orders_rows, orders_df,
                                    run_all, xbytes)

DISC = "Mira Goodfood Deal January"


def export(rows, invoice_date=None, catalog=None):
    cat = catalog if catalog is not None else catalog_from(HOT_BASIC, COLD_BASIC)[0]
    orders = load_orders_rows(rows)
    _, res, lines, _ = run_all(orders, cat)
    wb = openpyxl.load_workbook(io.BytesIO(eng.to_template_excel(res, lines, DISC, invoice_date)))
    return res, lines, wb


def blocks(ws):
    """{số thứ tự hóa đơn: [các dòng]} — bỏ qua dòng trống giữa các hóa đơn."""
    out = {}
    for row in ws.iter_rows(min_row=9, values_only=True):
        if all(v is None for v in row):
            continue
        out.setdefault(row[1], []).append(row)
    return out


def test_header_rows_match_template():
    _, _, wb = export([{"id": "A", "amount": 50_000, "discount": -5_000}])
    ws = wb["Hóa đơn bán hàng"]
    assert ws["B1"].value == eng.TEMPLATE_TITLE and ws["B1"].font.b
    assert [c.value for c in ws[8]][1:19] == eng.TEMPLATE_HEADERS
    assert ws["A8"].value is None
    assert wb.sheetnames == ["Hóa đơn bán hàng", "Kiểm tra"]


def test_block_structure_items_then_one_discount_then_total():
    res, lines, wb = export([{"id": "A", "short": "GM-1", "amount": 120_000, "discount": -12_345}])
    (rows,) = blocks(wb["Hóa đơn bán hàng"]).values()
    items, disc, total = rows[:-2], rows[-2], rows[-1]
    assert len(items) == len(lines)
    for r in items:
        assert r[13] * r[14] == r[15] and r[17] is None
    assert disc[11] == DISC and disc[15] == 12_345 and disc[17] == 1 and disc[13] is None
    assert total[11] == "THÀNH TIỀN GM-1" and total[15] == 120_000 - 12_345
    for r in rows:
        assert r[0] == "GM-1" and r[1] == 1 and r[6] == eng.BUYER and r[10] == eng.PAYMENT
        assert r[16] == "GRAB - GM-1"


def test_styles_discount_bold_total_filled_money_format():
    _, _, wb = export([{"id": "A", "amount": 120_000, "discount": -10_000}])
    ws = wb["Hóa đơn bán hàng"]
    last = max(r for r in range(9, ws.max_row + 1) if ws.cell(r, 12).value)
    assert ws.cell(last, 12).fill.fgColor.rgb == "FFD9EAF7" and ws.cell(last, 12).font.b
    assert ws.cell(last - 1, 12).font.b and ws.cell(last - 1, 16).number_format == "#,##0"
    assert ws.cell(9, 16).number_format == "#,##0" and ws.cell(9, 1).border.left.style == "thin"


def test_zero_discount_has_no_discount_line():
    _, _, wb = export([{"id": "A", "amount": 60_000, "discount": 0}])
    (rows,) = blocks(wb["Hóa đơn bán hàng"]).values()
    assert not any(r[17] == 1 for r in rows)
    assert rows[-1][15] == 60_000
    assert wb["Kiểm tra"]["F2"].value == 0 and wb["Kiểm tra"]["L2"].value == "OK"


def test_positive_discount_keeps_arithmetic():
    # giảm giá dương bất thường: dòng giảm giá mang số âm để THÀNH TIỀN vẫn = số tiền + giảm giá
    _, _, wb = export([{"id": "A", "amount": 60_000, "discount": 5_000}])
    (rows,) = blocks(wb["Hóa đơn bán hàng"]).values()
    assert rows[-2][15] == -5_000 and rows[-1][15] == 65_000


def test_invoices_sorted_by_created_date_and_numbered():
    df = orders_df([{"id": "late", "short": "GM-L", "amount": 50_000},
                    {"id": "early", "short": "GM-E", "amount": 60_000}])
    df["Ngày tạo"] = ["31 Jan 2026 9:17 PM", "14 Jan 2026 7:35 AM"]
    cat = catalog_from(HOT_BASIC, COLD_BASIC)[0]
    _, res, lines, _ = run_all(eng.load_orders(xbytes(df)), cat)
    wb = openpyxl.load_workbook(io.BytesIO(eng.to_template_excel(res, lines, DISC)))
    b = blocks(wb["Hóa đơn bán hàng"])
    assert b[1][0][0] == "GM-E" and b[1][0][2] == "14/01/2026"
    assert b[2][0][0] == "GM-L" and b[2][0][2] == "31/01/2026"


def test_fixed_invoice_date_applies_to_all():
    _, _, wb = export([{"id": "A", "amount": 50_000}, {"id": "B", "amount": 70_000}],
                      invoice_date=pd.Timestamp("2026-02-05"))
    dates = {r[2] for rows in blocks(wb["Hóa đơn bán hàng"]).values() for r in rows}
    assert dates == {"05/02/2026"}


def test_blank_row_between_invoices_only():
    _, _, wb = export([{"id": "A", "amount": 50_000}, {"id": "B", "amount": 70_000}])
    ws = wb["Hóa đơn bán hàng"]
    values = [[c.value for c in r] for r in ws.iter_rows(min_row=9)]
    blank = [i for i, r in enumerate(values) if all(v is None for v in r)]
    assert len(blank) == 1 and values[blank[0] - 1][11].startswith("THÀNH TIỀN")
    assert values[-1][11].startswith("THÀNH TIỀN")


def test_duplicate_short_ids_get_distinct_invoice_numbers():
    _, _, wb = export([{"id": "A", "short": "GM-9", "amount": 50_000},
                       {"id": "B", "short": "GM-9", "amount": 70_000}])
    b = blocks(wb["Hóa đơn bán hàng"])
    assert sorted(b) == [1, 2] and all(rows[0][0] == "GM-9" for rows in b.values())


def test_check_sheet_lists_failed_orders_and_flags_net_mismatch():
    rows = [{"id": "OK1", "amount": 50_000, "discount": -5_000},
            {"id": "BAD", "amount": 1_000, "discount": 0},  # nhỏ hơn món rẻ nhất -> không ghép được
            {"id": "LECH", "amount": 60_000, "discount": -6_000, "net": 50_000}]
    res, _, wb = export(rows)
    ks = {r[2]: r for r in wb["Kiểm tra"].iter_rows(min_row=2, values_only=True)}
    assert ks["OK1"][-1] == "OK"
    assert ks["BAD"][0] is None and ks["BAD"][-1].startswith("CHƯA GHÉP")
    assert ks["LECH"][-1].startswith("LỆCH") and ks["LECH"][10] == 54_000 and ks["LECH"][9] == 50_000
    assert len(blocks(wb["Hóa đơn bán hàng"])) == 2  # đơn lỗi không có trong sheet hóa đơn


def test_no_orders_exports_header_only():
    _, _, wb = export([{"id": "A", "status": "Đã hủy", "amount": 50_000}])
    ws = wb["Hóa đơn bán hàng"]
    assert ws.max_row == 8 and wb["Kiểm tra"].max_row == 1


@pytest.mark.parametrize("seed", [1, 7])
def test_every_block_reconciles_with_report(seed):
    rows = [{"id": f"O{i}", "amount": a, "discount": -d}
            for i, (a, d) in enumerate([(25_000, 2_500), (64_000, 0), (183_850, 18_385), (473_850, 47_000),
                                        (1_250_000, 125_000), (8_000, 800)])]
    cat = catalog_from(HOT_BASIC, COLD_BASIC)[0]
    _, res, lines, _ = run_all(load_orders_rows(rows), cat, seed=seed)
    wb = openpyxl.load_workbook(io.BytesIO(eng.to_template_excel(res, lines, DISC)))
    order = eng.invoice_order(res).set_index("invoice_no")
    for no, rws in blocks(wb["Hóa đơn bán hàng"]).items():
        items = sum(r[15] for r in rws if r[17] != 1 and not str(r[11]).startswith("THÀNH TIỀN"))
        disc = sum(r[15] for r in rws if r[17] == 1)
        assert items == int(order.loc[no, "amount"])
        assert rws[-1][15] == items - disc == int(order.loc[no, "net"])
